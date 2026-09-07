"""
Turn a tailored CV (markdown, as produced by llm.tailor_cv) into a clean
.docx file. Pure Python via python-docx — no system dependencies, so it runs
anywhere `pip install python-docx` works.

Handles the markdown subset the model produces for CVs:
  # Name            -> title
  ## Section        -> section heading
  - bullet          -> bulleted list item
  **bold** inline   -> bold runs
  plain line        -> normal paragraph
"""

import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

ACCENT = RGBColor(0x1F, 0x3A, 0x5F)   # dark blue for the name + headings
BOLD_SPLIT = re.compile(r"\*\*(.+?)\*\*")


def _add_inline(paragraph, text: str) -> None:
    """Add text to a paragraph, honouring **bold** segments."""
    pos = 0
    for m in BOLD_SPLIT.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        run = paragraph.add_run(m.group(1))
        run.bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def markdown_cv_to_docx(markdown_text: str, out_path: str) -> str:
    doc = Document()

    # Base document style.
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.6)
        section.left_margin = section.right_margin = Inches(0.7)

    for raw in markdown_text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue

        if line.startswith("# "):                      # name / title
            p = doc.add_paragraph()
            run = p.add_run(line[2:].strip())
            run.bold = True
            run.font.size = Pt(20)
            run.font.color.rgb = ACCENT
            p.space_after = Pt(2)

        elif line.startswith("## "):                    # section heading
            p = doc.add_paragraph()
            p.space_before = Pt(10)
            p.space_after = Pt(2)
            run = p.add_run(line[3:].strip().upper())
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = ACCENT

        elif line.startswith("- ") or line.startswith("* "):   # bullet
            p = doc.add_paragraph(style="List Bullet")
            _add_inline(p, line[2:].strip())

        else:                                           # normal paragraph
            p = doc.add_paragraph()
            _add_inline(p, line.strip())

    doc.save(out_path)
    return out_path
